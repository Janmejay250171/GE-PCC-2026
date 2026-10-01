from docx import Document
from docx.shared import Inches, Pt, RGBColor
from .styles import (
    add_heading_1, add_heading_2, add_heading_3,
    add_bullet, add_callout, format_table
)

def build_part2(doc: Document):
    # ---------------------------------------------------------
    # PHASE 6: SEHATSURE CORE INSURANCE & DECISION LOGIC
    # ---------------------------------------------------------
    add_heading_1(doc, "PHASE 6 — CORE INSURANCE MATHEMATICS & CLAUSE LOGIC")

    add_heading_2(doc, "6.1 Room Rent Capping & Proportionate Deduction Formula")
    add_bullet(doc, "The Healthcare Problem", "In India, insurance policies often cap daily room rent (e.g. ₹3,000/day or 1% of Sum Insured). When a patient chooses a room exceeding this cap (e.g. ₹6,000/day), insurers do NOT merely deduct the room difference; they invoke 'Proportionate Deduction', penalizing doctor fees, surgery charges, and nursing costs by the exact proportion of the room breach!")
    add_bullet(doc, "The Mathematical Formula Implemented",
               "AllowedRatio = Min(1.0, RoomLimitEligiblePerDay / RoomRatePerDay)\n"
               "AssociatedMedicalExpenses = ProcedureCharges (less sublimit excess) + DoctorFees\n"
               "ProportionateDisallowance = Round(AssociatedMedicalExpenses * (1.0 - AllowedRatio))\n"
               "Source: server/src/services/hospitalService.ts (lines 1188-1197)")
    add_bullet(doc, "Concrete Calculation Example",
               "Policy: ₹3,00,000 SI with 1% room cap (₹3,000/day). Selected Room: Single Private Room at ₹6,000/day (Stay: 2 days).\n"
               "Room Rent Excess = (₹6,000 - ₹3,000) * 2 = ₹6,000.\n"
               "AllowedRatio = ₹3,000 / ₹6,000 = 0.50 (50%).\n"
               "If Procedure = ₹50,000 and Doctor Fees = ₹10,000 -> Associated Expenses = ₹60,000.\n"
               "Proportionate Disallowance = ₹60,000 * (1 - 0.50) = ₹30,000 penalty!\n"
               "Patient pays ₹6,000 (room difference) + ₹30,000 (proportionate penalty) = ₹36,000 out-of-pocket before copay!")

    add_heading_2(doc, "6.2 Procedure-Specific Sub-Limits (Treatment Caps)")
    add_bullet(doc, "Implementation Logic", "If policy.subLimits contains a cap for a procedure (e.g. Knee Replacement capped at ₹1,50,000 or Cataract at ₹25,000):\n"
               "SublimitExcess = Max(0, ProcedureCharges - ApplicableSublimit)\n"
               "This excess is subtracted from the admissible claim and added directly to patient out-of-pocket.")

    add_heading_2(doc, "6.3 Compulsory Deductibles vs Co-Payments")
    add_bullet(doc, "Order of Application", "In health insurance accounting, deductions must follow strict legal sequence:\n"
               "1. Total Bill -> 2. Deduct Non-admissible / Room Excess / Proportionate Disallowance -> 3. Deduct Compulsory Deductible -> 4. Apply Co-Payment Percentage -> 5. Cap at Sum Insured -> 6. Add 5% Non-medical Consumables.")
    add_bullet(doc, "Safeguard for Missing Deductible", "If policy.deductible is null/undefined in the document, deductibleApplied is null and deductibleUnknown is true. The engine never arbitrarily invents a deductible.")

    add_heading_2(doc, "6.4 Modelled Non-Medical Expenses (5% IRDAI Consumables)")
    add_bullet(doc, "Regulatory Grounding", "IRDAI guidelines mandate that hospital consumables (PPE kits, gloves, sanitizers, admission kits, administrative fees) are non-reimbursable. Standard claims experience in India shows these average 5% to 8% of the total bill.")
    add_bullet(doc, "Formula", "ModelledNonMedicalAllowance = Round(TotalBill * 0.05). This is explicitly tagged as MODELLED ESTIMATE in the UI so the patient is financially prepared.")

    # ---------------------------------------------------------
    # PHASE 7: POLICY PDF PIPELINE & EXTRACTION TRUTH
    # ---------------------------------------------------------
    add_heading_1(doc, "PHASE 7 — POLICY PDF EXTRACTION PIPELINE")

    add_heading_2(doc, "7.1 Dual-Model LLM Ingestion Architecture")
    add_bullet(doc, "Primary & Fallback Models", "Primary: gemini-3.1-flash-lite. Fallback: gemini-3.5-flash-lite. Configured in server/src/config/env.ts with temperature=0 for zero randomness.")
    add_bullet(doc, "Hybrid Text & Vision Fallback", "geminiService.ts first attempts native text extraction using 'pdf-parse'. If extracted text > 50 characters, it passes text to Gemini. If text extraction yields <= 50 characters (indicating a scanned photocopy or image-only PDF), it automatically converts the PDF buffer to inline base64 and passes it as a visual multimodal document!")
    add_bullet(doc, "Transient Error Handling", "3-attempt exponential retry loop on HTTP 503, 429, timeout, or econnreset before invoking the fallback model.")

    add_heading_2(doc, "7.2 Self-Correcting Zod Validation Loop")
    add_bullet(doc, "Validation Engine", "ExtractedPolicyZod in server/src/schemas/policySchema.ts enforces types, enums, numbers, and structures.")
    add_bullet(doc, "Self-Correction Prompt", "If Zod validation fails, the service does not crash. It extracts the validation error messages and re-prompts Gemini:\n"
               "'IMPORTANT: Your previous response failed validation with these errors: [errors]. Return corrected JSON only.'\n"
               "This automated self-healing loop recovers from occasional formatting slips.")

    add_heading_2(doc, "7.3 Tier 1 Required Fields vs Tier 2 Defaults")
    add_bullet(doc, "Tier 1 Required Fields", "7 non-negotiable fields: insurer, policyType, sumInsured, roomLimit, copay, deductible, proportionateDeduction. Missing any blocks user confirmation unless resolved.")
    add_bullet(doc, "Tier 2 Defaults & Confidence", "Non-critical fields (ICU limit, sub-limits, ambulance limit, pre/post days) receive default values via applyTier2Defaults() and are marked with confidence='assumed'.")

    # ---------------------------------------------------------
    # PHASE 8: HOSPITAL DATASET FORENSIC AUDIT
    # ---------------------------------------------------------
    add_heading_1(doc, "PHASE 8 — HOSPITAL DATASET FORENSIC AUDIT")

    add_heading_2(doc, "8.1 Dataset Provenance & Streaming Ingestion")
    add_bullet(doc, "File & Records", "hospitals.csv at root: 11 MB containing 54,000+ facilities across India.")
    add_bullet(doc, "Columns", "hospital_name, hospital_type, address, insurers, rating, specialties, tier, segment.")
    add_bullet(doc, "Memory Optimization", "Instead of loading all 54k rows into JSON objects simultaneously, HospitalService uses readline streaming and internString() with STRING_POOL to reuse identical strings (e.g. 'Private', 'Bengaluru', 'Standard Private'). This keeps server memory under 150MB on Render.")
    add_bullet(doc, "Dataset Validation Reporting", "DatasetStats logs totalRecordsParsed, validRecordsLoaded, malformedRecordsRejected, missingInsurerCount, missingSpecialtyCount, missingAddressCount, and unknownCityCount.")

    add_heading_2(doc, "8.2 Address Parsing & Canonical City Extraction")
    add_bullet(doc, "The Challenge", "Indian addresses are notoriously unstructured (e.g. 'NARWAL BYE PASS ROAD JAMMU, Jammu and Kashmir, 180006').")
    add_bullet(doc, "Heuristic Cleaning Pipeline", "extractCity() splits address by commas, strips pincodes, discards state names (INDIAN_STATES set), strips road markers ('bypass road', 'national highway', 'marg', 'chowk'), and maps aliases ('bombay' -> 'Mumbai', 'bangalore' -> 'Bengaluru').")
    add_bullet(doc, "Canonical Tiers", "CITY_CANONICAL_TIERS maps over 80 major Indian cities into standardized tiers: Metro 1 (Delhi, Mumbai, Bengaluru, Chennai, Kolkata, Hyderabad), Metro 2 (Pune, Ahmedabad), Large City 1 (Jaipur, Lucknow), etc.")

    add_heading_2(doc, "8.3 4-State Network Matching & Patient Safeguards")
    headers_network = ["Network Status", "Match Condition", "UI Display / Action", "Cashless Eligibility"]
    rows_network = [
        ["verified", "Exact or canonical alias match in hospital empanelled insurers list", "Green Badge: 'Verified Network Partner'", "Eligible for Cashless Pre-Auth"],
        ["unverified", "Hospital record exists, but insurers list is empty or invalid", "Amber Badge: 'Network Status Not Verified'", "Requires check at TPA desk"],
        ["no_match", "Hospital lists insurers, but policy insurer is not listed", "Gray Badge: 'No Verified Network Match'", "Reimbursement Only"],
        ["unknown", "Policy insurer is unspecified or missing", "Blue Badge: 'Network Status Unknown'", "Requires policy clarification"]
    ]
    tbl_network = doc.add_table(rows=1, cols=4)
    format_table(tbl_network, [1.2, 2.5, 2.2, 1.3], headers_network, rows_network)

    add_callout(
        doc,
        "CRITICAL PATIENT SAFEGUARD: An unverified hospital is NEVER presented as network or cashless. "
        "Furthermore, networkMatchingService explicitly attaches freshness metadata: "
        "'Freshness: Unknown (Reference Dataset Match)'. We never mislead patients into false cashless assurances.",
        title="EMPANELED NETWORK INTEGRITY SAFEGUARD",
        alert_type="danger"
    )

    # ---------------------------------------------------------
    # PHASE 9: COST ESTIMATION & PRICING ENGINE
    # ---------------------------------------------------------
    add_heading_1(doc, "PHASE 9 — COST ESTIMATION & PRICING ENGINE")

    add_heading_2(doc, "9.1 Three-Tier Estimation Granularity")
    add_bullet(doc, "Level 1: Exact Procedure Match", "Matches Hospital Type + Tier + Specialty + Procedure in pvt_costs.csv / govt_costs.csv. Provides highest fidelity quote.")
    add_bullet(doc, "Level 2: Specialty Average", "If no specific procedure is selected, averages low, mean, high, and room rates across all procedures in that specialty for that tier.")
    add_bullet(doc, "Level 3: City / Tier Average", "If searching by city alone without specialty, averages all clinical procedures across the hospital's geographic tier.")

    add_heading_2(doc, "9.2 Segment Positioning Multipliers")
    add_bullet(doc, "Mathematical Curve", "SEGMENT_POSITION: Budget/Local (0.25), Standard Private (0.40), Established (0.50), Premium (0.65), Luxury (0.80), Government (0.50).\n"
               "If pos <= 0.50: Cost = low_cost + (mean_cost - low_cost) * (pos / 0.50)\n"
               "If pos > 0.50: Cost = mean_cost + (highest_cost - mean_cost) * ((pos - 0.50) / 0.50)\n"
               "This creates smooth, realistic pricing curves bounded strictly within dataset benchmarks.")

    add_heading_2(doc, "9.3 Deterministic Hash Variation (No Math.random)")
    add_bullet(doc, "The Algorithm", "getHospitalVariation() computes a hash from [hospital_name, address, tier, segment].join('|'). It scales the hash norm into a fixed offset strictly between -5% and +5%.\n"
               "Why? The same hospital always displays the exact same estimate across refreshes, avoiding random number flicker.")

    # ---------------------------------------------------------
    # PHASE 10: AI/ML TRUTH AUDIT
    # ---------------------------------------------------------
    add_heading_1(doc, "PHASE 10 — AI/ML TRUTH AUDIT (DEMARCATING REALITY)")

    add_callout(
        doc,
        "Never tell judges 'Everything is AI'. Technical judges immediately lose respect for teams that call basic if-else logic AI. "
        "Use this exact demarcation. It demonstrates architectural sophistication and intellectual honesty.",
        title="AI TRANSPARENCY DIRECTIVE",
        alert_type="warning"
    )

    headers_ai = ["Component / Feature", "True Classification", "Implementation Details", "Why This Choice Was Made"]
    rows_ai = [
        ["Policy Document Ingestion", "ACTUAL AI / LLM", "Gemini 3.1 Flash Lite with structured JSON prompt, temperature 0, and Zod validation", "LLMs excel at natural language comprehension of ambiguous, unstructured PDF contracts"],
        ["Statutory Scheme Rules", "RULE-BASED INTELLIGENCE", "Deterministic statutory overrides in overrideService.ts", "Statutory law must be 100% deterministic and legally auditable without hallucination"],
        ["Hospital Search & Ranking", "DATA-DRIVEN & MCDA", "In-memory streaming index + 4-pillar multi-attribute decision scoring formula", "Multi-criteria decision analysis provides transparent, explainable ranking weights"],
        ["Cost Estimation Engine", "DATA-DRIVEN DETERMINISTIC", "Segment positioning math + deterministic hash variation bounded by benchmark CSVs", "Clinical cost bounds must reflect empirical dataset baselines, not generative AI fiction"],
        ["Network Matching Service", "RULE-BASED DETERMINISTIC", "Canonical alias registry with token boundary regex matching", "Insurer empanelment requires exact legal alias verification to protect patient claims"],
        ["Care Journey Guidance", "EXPERT RULE ENGINE", "14 deterministic event handlers evaluating policy clauses & firing alerts", "Hospital guidance requires strict medical-insurance accuracy; AI cannot gamble with disallowances"],
        ["Natural Language Explanations", "AI / TEMPLATE HYBRID", "Synthesizes calculations into structured 'Why Am I Seeing This?' explainability cards", "Separates mathematical calculation from natural language communication"]
    ]
    tbl_ai = doc.add_table(rows=1, cols=4)
    format_table(tbl_ai, [1.5, 1.4, 2.3, 2.0], headers_ai, rows_ai)
