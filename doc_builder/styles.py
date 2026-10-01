import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Color Palette Constants
COLOR_PRIMARY_HEX = "0F2942"      # Deep Navy / Slate
COLOR_SECONDARY_HEX = "1D68BD"    # Medical Blue
COLOR_ACCENT_HEX = "0284C7"       # Sky Blue
COLOR_SUCCESS_HEX = "15803D"      # Dark Green
COLOR_WARNING_HEX = "B45309"      # Dark Amber
COLOR_DANGER_HEX = "B91C1C"       # Dark Crimson
COLOR_MUTED_HEX = "64748B"        # Slate Gray
COLOR_BG_LIGHT_HEX = "F8FAFC"     # Off-white / light slate
COLOR_BG_ALERT_HEX = "FEF2F2"     # Light red
COLOR_BG_WARN_HEX = "FFFBEB"      # Light amber
COLOR_BG_SUCCESS_HEX = "F0FDF4"   # Light green
COLOR_BG_INFO_HEX = "EFF6FF"      # Light blue

COLOR_PRIMARY = RGBColor(15, 41, 66)
COLOR_SECONDARY = RGBColor(29, 104, 189)
COLOR_TEXT = RGBColor(30, 41, 59)
COLOR_MUTED = RGBColor(100, 116, 139)
COLOR_DANGER = RGBColor(185, 28, 28)
COLOR_SUCCESS = RGBColor(21, 128, 61)
COLOR_WARNING = RGBColor(180, 83, 9)

def setup_document_styles(doc: Document):
    """Sets page margins and baseline typography styles."""
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Base Normal Style
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(10.5)
    font.color.rgb = COLOR_TEXT
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(4)

def add_header_footer(doc: Document):
    for s_idx, section in enumerate(doc.sections):
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("SEHATSURE — GE Precision Care Challenge 2026 | Final Round Defense Manual")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_MUTED

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("CONFIDENTIAL & AUTHORITATIVE — For Candidate Defense & Live Demo Preparation Only")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8)
        frun.font.color.rgb = COLOR_MUTED

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def add_title(doc: Document, text: str, subtitle: str = ""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY

    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(16)
        r2 = p2.add_run(subtitle)
        r2.font.name = 'Calibri'
        r2.font.size = Pt(13)
        r2.font.color.rgb = COLOR_SECONDARY
        r2.font.bold = True

def add_heading_1(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(17)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p

def add_heading_2(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(13.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY
    return p

def add_heading_3(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p

def add_bullet(doc: Document, title: str, text: str, bold_title: bool = True):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if title:
        rt = p.add_run(title + (": " if bold_title else " "))
        rt.font.name = 'Calibri'
        rt.font.bold = bold_title
        rt.font.color.rgb = COLOR_PRIMARY
    rx = p.add_run(text)
    rx.font.name = 'Calibri'
    rx.font.color.rgb = COLOR_TEXT
    return p

def add_callout(doc: Document, text: str, title: str = "IMPORTANT NOTICE", alert_type: str = "info"):
    bg_map = {
        "info": (COLOR_BG_INFO_HEX, COLOR_SECONDARY_HEX, COLOR_SECONDARY),
        "warning": (COLOR_BG_WARN_HEX, COLOR_WARNING_HEX, COLOR_WARNING),
        "danger": (COLOR_BG_ALERT_HEX, COLOR_DANGER_HEX, COLOR_DANGER),
        "success": (COLOR_BG_SUCCESS_HEX, COLOR_SUCCESS_HEX, COLOR_SUCCESS)
    }
    bg_hex, border_hex, title_color = bg_map.get(alert_type, bg_map["info"])

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    # Left thick border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
                        f'<w:top w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'<w:bottom w:val="none"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(10.5)
    r_title.font.color.rgb = title_color

    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(10.5)
    r_body.font.color.rgb = COLOR_TEXT

    # Add spacing after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def format_table(tbl, col_widths, headers, rows):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Format Header Row
    hdr_cells = tbl.rows[0].cells
    for i, h_text in enumerate(headers):
        cell = hdr_cells[i]
        set_cell_background(cell, COLOR_PRIMARY_HEX)
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(h_text)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Populate and Format Data Rows
    for r_idx, row_data in enumerate(rows):
        row_cells = tbl.add_row().cells
        bg_color = COLOR_BG_LIGHT_HEX if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            cell = row_cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(str(cell_value))
            run.font.name = 'Calibri'
            run.font.size = Pt(9)
            run.font.color.rgb = COLOR_TEXT

    # Set Column Widths
    for row in tbl.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)

    # Set subtle borders
    for row in tbl.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                                f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                                f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                                f'<w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                                f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                                f'</w:tcBorders>')
            tcPr.append(borders)
