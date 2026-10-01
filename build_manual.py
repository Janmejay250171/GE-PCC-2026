import os
import sys
import docx
from docx import Document
from doc_builder.styles import setup_document_styles, add_header_footer
from doc_builder.sections_part1 import build_part1
from doc_builder.sections_part2 import build_part2
from doc_builder.sections_part3 import build_part3
from doc_builder.sections_part4 import build_part4
from doc_builder.sections_part5 import build_part5
from doc_builder.sections_part6 import build_part6

def main():
    print("=" * 60)
    print("SEHATSURE — GENERATING FINAL ROUND DEFENSE & DEMO MANUAL")
    print("GE HealthCare Precision Care Challenge 2026")
    print("=" * 60)

    doc = Document()
    setup_document_styles(doc)

    print("-> Assembling Part 1: Executive Overview & Phases 1-5...")
    build_part1(doc)

    print("-> Assembling Part 2: Core Insurance Mathematics & Phases 6-10...")
    build_part2(doc)

    print("-> Assembling Part 3: Architecture, Backend, Frontend & Phases 11-15...")
    build_part3(doc)

    print("-> Assembling Part 4: Weakness Defense, 100+ Q&As & Phases 16-20...")
    build_part4(doc)

    print("-> Assembling Part 5: 10-Minute Script, Pitches, Trade-offs & Phases 21-25...")
    build_part5(doc)

    print("-> Assembling Part 6: Provenance, Cheat Sheet, 4h Study Guide & Phases 26-35...")
    build_part6(doc)

    add_header_footer(doc)

    output_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "SEHATSURE_FINAL_ROUND_DEFENSE_AND_DEMO_MANUAL.docx"))
    doc.save(output_path)

    file_size_kb = os.path.getsize(output_path) / 1024
    para_count = len(doc.paragraphs)
    table_count = len(doc.tables)

    print("=" * 60)
    print(f"SUCCESS: Document generated at:")
    print(f" -> {output_path}")
    print(f"File Size: {file_size_kb:.2f} KB")
    print(f"Total Paragraphs: {para_count}")
    print(f"Total Tables: {table_count}")
    print("=" * 60)

if __name__ == "__main__":
    main()
